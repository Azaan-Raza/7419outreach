"""Builds excel/OutreachHours.xlsx — the spreadsheet version of the tracker.

Run:  python3 excel/build_workbook.py   (needs openpyxl)
Sheets: Start here, Dashboard, Members, Sessions, Settings, Sign up, Log hours.
"""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

NAVY, NAVY2, GOLD, GOLD_DARK, CREAM, RED, GRAY = "11224E", "1A2F5E", "FFC14A", "926408", "F5F5EE", "DC2626", "8A8575"
MEMBER_ROWS, SESSION_ROWS = 300, 1000

wb = Workbook()
title_font = Font(name="Georgia", size=20, bold=True, color=NAVY)
sub_font = Font(name="Calibri", size=11, color=GRAY, italic=True)
body = Font(name="Calibri", size=11, color=NAVY)
head_font = Font(name="Calibri", size=11, bold=True, color=GOLD)
head_fill = PatternFill("solid", fgColor=NAVY)
label_font = Font(name="Calibri", size=11, bold=True, color=NAVY)
input_fill = PatternFill("solid", fgColor="FFFFFF")
thin = Side(style="thin", color="D9D9CF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, cols, row=1):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = title_font
    if sub:
        ws["A2"] = sub
        ws["A2"].font = sub_font
    ws.sheet_view.showGridLines = False


def name(nm, ref):
    dn = DefinedName(nm, attr_text=ref)
    try:
        wb.defined_names[nm] = dn
    except TypeError:
        wb.defined_names.append(dn)


# ---------------------------------------------------------------- Start here
ws = wb.active
ws.title = "Start here"
ws.sheet_properties.tabColor = NAVY
title(ws, "Outreach Tracker — Team 7419", "The spreadsheet version of the tracker. Everything counts toward the semester in Settings.")
lines = [
    ("How it works", True),
    ("1. Members sign up once on the Sign up sheet (or a lead adds them straight into Members).", False),
    ("2. Each outreach shift goes on the Sessions sheet: date, member email, event, clock in, clock out, and links to the two photos.", False),
    ("3. A lead sets Status to Approved (and can override the hours) or Sent back with a note.", False),
    ("4. Members shows approved hours, hours waiting on review, and a progress bar against the required hours.", False),
    ("", False),
    ("Scripts", True),
    ("Excel on the web / Microsoft 365: Automate → New Script → paste OutreachTracker.ts → save. Add a button on the Sign up and Log hours sheets that runs it.", False),
    ("Desktop Excel: import OutreachTracker.bas (Developer → Visual Basic → File → Import) and save the workbook as .xlsm. Macros: SignUp, LogSession, RunForms, ApproveSelected, SendBackSelected, ShowProgress.", False),
    ("", False),
    ("From the web app", True),
    ("Leads can download all sessions as a .csv from the Members tab. Its columns line up with the Sessions sheet, so you can paste the rows straight in.", False),
]
for i, (t, bold) in enumerate(lines, 4):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name="Calibri", size=12 if bold else 11, bold=bold, color=NAVY)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120

# ---------------------------------------------------------------- Settings
st = wb.create_sheet("Settings")
st.sheet_properties.tabColor = GOLD
title(st, "Settings", "Only sessions between these dates count.")
rows = [("Semester", "Fall 2026"), ("Starts", date(2026, 8, 17)), ("Ends", date(2026, 12, 18)), ("Hours each member needs", 9), ("What has to be in every photo", "the object specified by a lead")]
for i, (k, v) in enumerate(rows, 4):
    st.cell(row=i, column=1, value=k).font = label_font
    c = st.cell(row=i, column=2, value=v)
    c.font, c.fill, c.border = body, input_fill, box
    if isinstance(v, date):
        c.number_format = "mmm d, yyyy"
widths(st, [34, 28])
name("SemesterName", "Settings!$B$4")
name("SemesterStart", "Settings!$B$5")
name("SemesterEnd", "Settings!$B$6")
name("RequiredHours", "Settings!$B$7")
name("PhotoObject", "Settings!$B$8")

# ---------------------------------------------------------------- Members
mem = wb.create_sheet("Members")
mem.sheet_properties.tabColor = NAVY
header(mem, ["Name", "Email", "Grade", "Signed up", "Approved hours", "Waiting on review", "Progress", "Status", "Hours to go"])
for r in range(2, MEMBER_ROWS + 2):
    mem[f"E{r}"] = f'=IF($B{r}="","",SUMIFS(Sessions!$I:$I,Sessions!$B:$B,$B{r},Sessions!$A:$A,">="&SemesterStart,Sessions!$A:$A,"<="&SemesterEnd))'
    mem[f"F{r}"] = f'=IF($B{r}="","",SUMIFS(Sessions!$F:$F,Sessions!$B:$B,$B{r},Sessions!$G:$G,"Waiting",Sessions!$A:$A,">="&SemesterStart,Sessions!$A:$A,"<="&SemesterEnd))'
    mem[f"G{r}"] = f'=IF($B{r}="","",MIN(1,E{r}/RequiredHours))'
    mem[f"H{r}"] = f'=IF($B{r}="","",IF(E{r}>=RequiredHours,"Done","In progress"))'
    mem[f"I{r}"] = f'=IF($B{r}="","",MAX(0,RequiredHours-E{r}))'
    mem[f"D{r}"].number_format = "mmm d, yyyy"
    for col in "EFI":
        mem[f"{col}{r}"].number_format = "0.00"
    mem[f"G{r}"].number_format = "0%"
    for col in "ABCDEFGHI":
        mem[f"{col}{r}"].font = body
widths(mem, [24, 30, 8, 14, 15, 17, 16, 13, 12])
mem.freeze_panes = "A2"
mem.conditional_formatting.add(f"G2:G{MEMBER_ROWS + 1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=GOLD, showValue=True))
mem.conditional_formatting.add(f"H2:H{MEMBER_ROWS + 1}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor=NAVY), font=Font(color=GOLD, bold=True)))
dv_grade = DataValidation(type="list", formula1='"9,10,11,12"', allow_blank=True)
mem.add_data_validation(dv_grade)
dv_grade.add(f"C2:C{MEMBER_ROWS + 1}")

# ---------------------------------------------------------------- Sessions
ses = wb.create_sheet("Sessions")
ses.sheet_properties.tabColor = GOLD
header(ses, ["Date", "Member email", "Event", "Clock in", "Clock out", "Hours on clock", "Status", "Approved hours", "Counted hours", "Lead note", "Clock-in photo", "Clock-out photo"])
for r in range(2, SESSION_ROWS + 2):
    ses[f"F{r}"] = f'=IF(OR($D{r}="",$E{r}=""),"",ROUND(MOD($E{r}-$D{r},1)*24,2))'
    ses[f"I{r}"] = f'=IF($G{r}="Approved",IF($H{r}="",IF($F{r}="",0,$F{r}),$H{r}),0)'
    ses[f"A{r}"].number_format = "mmm d, yyyy"
    ses[f"D{r}"].number_format = "h:mm AM/PM"
    ses[f"E{r}"].number_format = "h:mm AM/PM"
    for col in "FHI":
        ses[f"{col}{r}"].number_format = "0.00"
    for col in "ABCDEFGHIJKL":
        ses[f"{col}{r}"].font = body
widths(ses, [13, 30, 28, 12, 12, 14, 12, 14, 13, 40, 22, 22])
ses.freeze_panes = "A2"
dv_status = DataValidation(type="list", formula1='"Waiting,Approved,Sent back"', allow_blank=True)
ses.add_data_validation(dv_status)
dv_status.add(f"G2:G{SESSION_ROWS + 1}")
dv_email = DataValidation(type="list", formula1=f"=Members!$B$2:$B${MEMBER_ROWS + 1}", allow_blank=True, showErrorMessage=False)
ses.add_data_validation(dv_email)
dv_email.add(f"B2:B{SESSION_ROWS + 1}")
ses.conditional_formatting.add(f"G2:G{SESSION_ROWS + 1}", CellIsRule(operator="equal", formula=['"Approved"'], fill=PatternFill("solid", fgColor=GOLD), font=Font(color=NAVY, bold=True)))
ses.conditional_formatting.add(f"G2:G{SESSION_ROWS + 1}", CellIsRule(operator="equal", formula=['"Sent back"'], font=Font(color=RED, bold=True)))
ses.conditional_formatting.add(f"G2:G{SESSION_ROWS + 1}", CellIsRule(operator="equal", formula=['"Waiting"'], font=Font(color=GRAY, italic=True)))
# sample rows so the formulas have something to show
samples = [
    (date(2026, 9, 3), "alex@example.com", "Library STEM night", "17:02", "19:35", "Approved", 2.5, "", "", ""),
    (date(2026, 9, 10), "alex@example.com", "Elementary school robot demo", "13:58", "16:01", "Approved", "", "", "", ""),
    (date(2026, 9, 14), "alex@example.com", "Farmers market booth", "9:04", "11:40", "Waiting", "", "", "", ""),
    (date(2026, 9, 14), "sam@example.com", "Farmers market booth", "9:30", "9:41", "Sent back", "", "Photos are 11 minutes apart. Log the full shift next time.", "", ""),
]
from datetime import datetime, time
for i, (d, em, ev, tin, tout, status, hrs, note, p1, p2) in enumerate(samples, 2):
    ses[f"A{i}"] = d
    ses[f"B{i}"] = em
    ses[f"C{i}"] = ev
    ses[f"D{i}"] = datetime.strptime(tin, "%H:%M").time()
    ses[f"E{i}"] = datetime.strptime(tout, "%H:%M").time()
    ses[f"G{i}"] = status
    if hrs != "":
        ses[f"H{i}"] = hrs
    ses[f"J{i}"] = note
for i, (n, em, g) in enumerate([("Alex Kim", "alex@example.com", "11"), ("Jordan Patel", "jordan@example.com", "10"), ("Sam Rivera", "sam@example.com", "12")], 2):
    mem[f"A{i}"], mem[f"B{i}"], mem[f"C{i}"], mem[f"D{i}"] = n, em, g, date(2026, 8, 20)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 1)
db.sheet_properties.tabColor = GOLD
title(db, "Outreach Tracker")
db["A2"] = "=SemesterName"
db["A2"].font = sub_font
db["A3"] = "Run the script or open Members for the live numbers."
db["A3"].font = sub_font
tiles = [("Members", "=COUNTA(Members!B2:B%d)" % (MEMBER_ROWS + 1)), ("Done", '=COUNTIF(Members!H:H,"Done")'), ("Hours approved", "=SUM(Members!E2:E%d)" % (MEMBER_ROWS + 1)),
         ("Sessions waiting", '=COUNTIF(Sessions!G:G,"Waiting")'), ("Hours still owed", "=SUM(Members!I2:I%d)" % (MEMBER_ROWS + 1))]
for i, (lbl, f) in enumerate(tiles):
    col = get_column_letter(i + 1)
    v = db[f"{col}5"]
    v.value = f
    v.font = Font(name="Georgia", size=24, bold=True, color=NAVY)
    v.fill = PatternFill("solid", fgColor=CREAM)
    v.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    v.number_format = "0.##"
    l = db[f"{col}6"]
    l.value = lbl.upper()
    l.font = Font(name="Calibri", size=9, bold=True, color=GRAY)
    l.fill = PatternFill("solid", fgColor=CREAM)
    l.alignment = Alignment(horizontal="left", indent=1)
db.row_dimensions[5].height = 40
db["A8"] = "Progress by member"
db["A8"].font = Font(name="Georgia", size=14, bold=True, color=NAVY)
header(db, ["Member", "Approved", "Progress", "Status", "To go"], row=9)
for k in range(60):
    r, m = 10 + k, 2 + k
    db[f"A{r}"] = f'=IF(Members!B{m}="","",Members!A{m})'
    db[f"B{r}"] = f'=IF(Members!B{m}="","",Members!E{m})'
    db[f"C{r}"] = f'=IF(Members!B{m}="","",REPT("█",ROUND(Members!G{m}*20,0))&REPT("░",20-ROUND(Members!G{m}*20,0)))'
    db[f"D{r}"] = f'=IF(Members!B{m}="","",Members!H{m})'
    db[f"E{r}"] = f'=IF(Members!B{m}="","",Members!I{m})'
    db[f"B{r}"].number_format = db[f"E{r}"].number_format = "0.00"
    db[f"C{r}"].font = Font(name="Consolas", size=11, color=GOLD_DARK)
    for col in "ABDE":
        db[f"{col}{r}"].font = body
widths(db, [26, 14, 30, 14, 12])
db.conditional_formatting.add("D10:D69", CellIsRule(operator="equal", formula=['"Done"'], font=Font(color=GOLD_DARK, bold=True)))

# ---------------------------------------------------------------- Sign up form
su = wb.create_sheet("Sign up")
su.sheet_properties.tabColor = NAVY2
title(su, "Sign up", "Fill in the three boxes, then run the Outreach tracker script (or the RunForms macro). Your row moves to Members and this clears.")
for i, k in enumerate(["Name", "Email", "Grade"], 4):
    su.cell(row=i, column=1, value=k).font = label_font
    c = su.cell(row=i, column=2)
    c.fill, c.border, c.font = input_fill, box, body
dv = DataValidation(type="list", formula1='"9,10,11,12"', allow_blank=True)
su.add_data_validation(dv)
dv.add("B6")
su["A8"] = "=\"Every member needs \"&RequiredHours&\" approved hours in \"&SemesterName&\".\""
su["A8"].font = sub_font
widths(su, [16, 40])

# ---------------------------------------------------------------- Log hours form
lg = wb.create_sheet("Log hours")
lg.sheet_properties.tabColor = NAVY2
title(lg, "Log hours", "One shift at a time. Fill it in, run the script (or RunForms macro), and it lands on Sessions as Waiting.")
labels = ["Your email", "Date", "Event", "Clock in", "Clock out", "Clock-in photo link", "Clock-out photo link"]
for i, k in enumerate(labels, 4):
    lg.cell(row=i, column=1, value=k).font = label_font
    c = lg.cell(row=i, column=2)
    c.fill, c.border, c.font = input_fill, box, body
lg["B5"].number_format = "mmm d, yyyy"
lg["B7"].number_format = lg["B8"].number_format = "h:mm AM/PM"
dv2 = DataValidation(type="list", formula1=f"=Members!$B$2:$B${MEMBER_ROWS + 1}", allow_blank=True, showErrorMessage=False)
lg.add_data_validation(dv2)
dv2.add("B4")
lg["A12"] = '="Both photos need to show "&PhotoObject&"."'
lg["A12"].font = sub_font
lg["A13"] = "Type times like 3:42 PM. Paste photo links from the app, Drive or Photos."
lg["A13"].font = sub_font
widths(lg, [22, 44])

wb.save("excel/OutreachHours.xlsx")
print("wrote excel/OutreachHours.xlsx")
